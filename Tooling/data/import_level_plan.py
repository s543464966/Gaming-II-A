"""只读提取新版装量分组与难度目标；独立复算数量，不把表内待验证字段当作结果。"""
from pathlib import Path
from collections import defaultdict
import argparse
import hashlib
import json
import re
import openpyxl


def extract(source):
    workbook = openpyxl.load_workbook(source, data_only=True)
    is_r2 = 'R2' in str(workbook['关卡总表']['A3'].value)
    adjustment_path = Path(__file__).with_name('level_adjustments.json')
    adjustments = json.loads(adjustment_path.read_text(encoding='utf-8'))
    groups = defaultdict(list)
    authored = {}
    for row in workbook['装量配置'].iter_rows(min_row=6, values_only=True):
        if not isinstance(row[0], int):
            continue
        groups[row[0]].append(dict(zip(
            ('level', 'phase', 'id', 'kind', 'count', 'capacity', 'fill', 'number'), row[:8])))
        # 手工参考关的层序仍由数值源维护，重复导入不得随机改写。
        if isinstance(row[16], str) and row[16].startswith('层序JSON:'):
            assert row[0] not in authored, ('Duplicate authored opening', row[0])
            authored[row[0]] = json.loads(row[16].removeprefix('层序JSON:'))
    levels = []
    for row in workbook['关卡总表'].iter_rows(min_row=6, values_only=True):
        if not isinstance(row[0], int):
            continue
        names = {0:'id', 2:'role', 3:'difficulty', 6:'layout', 7:'flavors', 8:'filled',
                 9:'stock', 10:'orders', 11:'empty', 12:'single', 14:'slots', 15:'general_space',
                 16:'interleave', 17:'hidden_ratio', 18:'aabb_max', 19:'demand_pressure',
                 20:'stock_pressure', 21:'main', 22:'secondary', 23:'parameters', 24:'donuts',
                 29:'baseline_rate_target', 30:'duration_target_minutes', 36:'composition', 37:'assist', 38:'assist_rate_target',
                 43:'opening_donuts', 44:'normal_fill_ratio', 45:'matching_space',
                 46:'restricted_space', 47:'locked_space', 48:'adjacent_max',
                 49:'first_group_min', 50:'one_step_max', 51:'special_count'}
        plan = {key: row[index] for index, key in names.items()}
        if is_r2:
            plan.update(special_types=row[52], prerequisite_boxes=row[53],
                        numeric_decrements=row[54], mechanism_target=row[55],
                        midgame_space_min=row[56], tuning_notes=row[35])
            multiplier = re.search(r'拆弹倍率([\d.]+)', str(plan['parameters']))
            plan['bomb_multiplier'] = float(multiplier.group(1)) if multiplier else None
        plan['source_row'] = plan['id'] + 5
        plan['groups'] = groups[plan['id']]
        if plan['id'] in authored:
            plan['authored_opening'] = authored[plan['id']]
        # R2 已重新配平并更换组 ID，旧版 48/49 的修正不能叠加到新表。
        plan['approved_adjustments'] = [] if is_r2 else adjustments['levels'].get(str(plan['id']), [])
        initial = [g for g in plan['groups'] if g['phase'] == '开局']
        stock = [g for g in plan['groups'] if g['phase'] == '备货']
        total = lambda items: sum(g['count'] * g['fill'] for g in items)
        assert sum(g['count'] for g in initial) == plan['slots'] <= 25, plan['id']
        assert total(initial) == plan['opening_donuts'], plan['id']
        assert total(stock) + total(initial) == plan['donuts'] == plan['orders'] * 4, plan['id']
        assert sum(g['count'] for g in stock) == plan['stock'], plan['id']
        assert sum(g['count'] for g in initial if g['kind'] == '广告锁定周转盒') == 2
        for group in plan['groups']:
            kind, fill = group['kind'], group['fill']
            allowed = range(5) if kind == '普通盒' else ((0,) if kind in ('固定颜色盒','广告锁定周转盒') else ((0,1) if kind in ('单向收纳盒','单颗暂存盒') else (3,4)))
            assert fill in allowed and fill <= group['capacity'], group
        effective = [dict(g, count=1) for g in plan['groups'] for _ in range(g['count'])]
        for change in plan['approved_adjustments']:
            if 'group' in change:
                group = next(g for g in effective if g['id'] == change['group'])
                group.update(kind=change['kind'], capacity=change['capacity'])
            else:
                next(g for g in effective if g['id']==change['from_group'])['fill'] -= change['amount']
                next(g for g in effective if g['id']==change['to_group'])['fill'] += change['amount']
        if plan['approved_adjustments']:
            active = [g for g in effective if g['phase']=='开局' and g['kind']!='广告锁定周转盒']
            normals = [g for g in active if g['kind']=='普通盒' and g['fill']>0]
            updates = {'filled':sum(g['capacity']==4 and g['fill']>0 for g in active),
                       'single':sum(g['capacity']==1 for g in active),
                       'special_count':sum(g['kind']!='普通盒' for g in active),
                       'normal_fill_ratio':sum(g['fill'] for g in normals)/(4*len(normals)),
                       'matching_space':sum(g['capacity']-g['fill'] for g in active if g['fill'] and g['kind'] in ('普通盒','单颗暂存盒','底部置顶盒','炸弹盒')),
                       'locked_space':sum(g['capacity']-g['fill'] for g in active if g['kind'] in ('数字盖盒','普通冰冻盒','数字冰冻盒'))}
            plan['source_values']={key:plan[key] for key in updates}
            plan.update(updates)
        plan['effective_groups'] = effective
        if is_r2:
            active = [g for g in effective if g['phase']=='开局' and g['kind']!='广告锁定周转盒']
            special = [g for g in active if g['kind']!='普通盒']
            normals = [g for g in active if g['kind']=='普通盒' and g['fill']>0]
            derived = {
                'filled':sum(g['capacity']==4 and g['fill']>0 for g in active),
                'empty':sum(g['capacity']==4 and g['fill']==0 for g in active),
                'single':sum(g['capacity']==1 for g in active),
                'general_space':sum(g['capacity'] for g in active if g['fill']==0 and g['kind'] in ('普通盒','单颗暂存盒')),
                'matching_space':sum(g['capacity']-g['fill'] for g in active if g['fill']>0 and g['kind'] in ('普通盒','单颗暂存盒','底部置顶盒','炸弹盒')),
                'restricted_space':sum(g['capacity']-g['fill'] for g in active if g['kind'] in ('固定颜色盒','单向收纳盒')),
                'locked_space':sum(g['capacity']-g['fill'] for g in active if g['kind'] in ('数字盖盒','普通冰冻盒','数字冰冻盒')),
                'prerequisite_boxes':sum(g['kind'] in ('数字盖盒','普通冰冻盒','数字冰冻盒') for g in active),
                'normal_fill_ratio':sum(g['fill'] for g in normals)/(4*len(normals)),
            }
            for key, value in derived.items():
                assert abs(value-plan[key]) < 1e-8, (plan['id'],key,value,plan[key])
            assert len(special) == plan['special_count'], plan['id']
            assert len({g['kind'] for g in special}) == plan['special_types'], plan['id']
            assert sum(g['number'] for g in active) == plan['numeric_decrements'], plan['id']
            assert plan['general_space'] + plan['matching_space'] >= 4, plan['id']
        levels.append(plan)
    assert [p['id'] for p in levels] == list(range(1,101))
    assert sum(p['stock'] > 0 for p in levels) == 30
    return {'source': str(source), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
            'sheet': '关卡总表', 'groups_sheet': '装量配置', 'requirements': 'R2' if is_r2 else 'V1.3',
            'layer_order': 'runtime_top_to_bottom', 'levels': levels}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    data = extract(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(f"PASS: {data['requirements']} 100 rows, {sum(len(p['groups']) for p in data['levels'])} groups, initial capacities and total conservation; 30 stock levels")
